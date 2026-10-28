"""
Commercial-Grade CSV (UTF-8-SIG BOM) & Styled Excel (.xlsx) Exporters.
Highlights NEW (soft green) and UPDATED (soft amber) rows for instant client review.
"""
from __future__ import annotations
import csv
from pathlib import Path
from typing import List

import openpyxl
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

from .diff_store import DiffRecord, DiffStatus


HEADERS_BILINGUAL = [
    "Diff Status (異動狀態)",
    "Record ID (編號)",
    "Region (縣市/地區)",
    "Organization (機關/企業名稱)",
    "Department / Duty (科室與業務職掌)",
    "Phone (代表電話)",
    "Extension (分機)",
    "Fax (傳真)",
    "Address (地址)",
    "Email / Website (信箱或網址)",
    "Source URL (來源網址)",
    "Scraped At (擷取時間 UTC+8)",
]


def export_diff_to_csv(diff_records: List[DiffRecord], csv_path: Path) -> Path:
    """Writes UTF-8 with BOM (`utf-8-sig`) CSV so Windows Excel opens Traditional Chinese cleanly."""
    csv_path.parent.mkdir(parents=True, exist_ok=True)
    with csv_path.open("w", encoding="utf-8-sig", newline="") as f:
        writer = csv.writer(f)
        writer.writerow(HEADERS_BILINGUAL)
        for item in diff_records:
            r = item.record
            writer.writerow(
                [
                    item.status.value,
                    r.record_id,
                    r.region,
                    r.organization,
                    r.department_or_duty,
                    r.phone,
                    r.extension,
                    r.fax,
                    r.address,
                    r.email_or_url,
                    r.source_url,
                    r.scraped_at,
                ]
            )
    return csv_path


def export_diff_to_excel(diff_records: List[DiffRecord], xlsx_path: Path) -> Path:
    """Exports a formatted `.xlsx` workbook with frozen header, auto-filter, and diff row highlighting."""
    xlsx_path.parent.mkdir(parents=True, exist_ok=True)
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Directory_Diff_Report"

    header_fill = PatternFill(start_color="1F4E78", end_color="1F4E78", fill_type="solid")
    header_font = Font(name="Calibri", size=11, bold=True, color="FFFFFF")
    new_fill = PatternFill(start_color="E6F4EA", end_color="E6F4EA", fill_type="solid")
    upd_fill = PatternFill(start_color="FEF7E0", end_color="FEF7E0", fill_type="solid")
    unchanged_fill = PatternFill(start_color="FFFFFF", end_color="FFFFFF", fill_type="solid")
    thin = Side(border_style="thin", color="D0D7DE")
    border = Border(left=thin, right=thin, top=thin, bottom=thin)

    ws.append(HEADERS_BILINGUAL)
    ws.row_dimensions[1].height = 26
    for col_idx in range(1, len(HEADERS_BILINGUAL) + 1):
        c = ws.cell(row=1, column=col_idx)
        c.fill = header_fill
        c.font = header_font
        c.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        c.border = border

    for row_idx, item in enumerate(diff_records, start=2):
        r = item.record
        vals = [
            item.status.value,
            r.record_id,
            r.region,
            r.organization,
            r.department_or_duty,
            r.phone,
            r.extension,
            r.fax,
            r.address,
            r.email_or_url,
            r.source_url,
            r.scraped_at,
        ]
        ws.append(vals)
        ws.row_dimensions[row_idx].height = 24
        row_fill = (
            new_fill
            if item.status == DiffStatus.NEW
            else (upd_fill if item.status == DiffStatus.UPDATED else unchanged_fill)
        )
        for col_idx in range(1, len(vals) + 1):
            c = ws.cell(row=row_idx, column=col_idx)
            c.fill = row_fill
            c.border = border
            c.alignment = Alignment(horizontal="left", vertical="center", wrap_text=True)

    ws.freeze_panes = "A2"
    ws.auto_filter.ref = f"A1:{get_column_letter(len(HEADERS_BILINGUAL))}{len(diff_records) + 1}"
    widths = [18, 14, 14, 28, 38, 18, 14, 18, 34, 32, 34, 22]
    for idx, w in enumerate(widths, start=1):
        ws.column_dimensions[get_column_letter(idx)].width = w

    wb.save(xlsx_path)
    return xlsx_path
