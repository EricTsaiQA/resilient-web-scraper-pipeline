"""
Comprehensive Pytest Suite for resilient-web-scraper-pipeline.
Tests:
1. Full-width Unicode normalization and phone/extension splitting
2. Deterministic SHA-256 diff detection (NEW, UPDATED, UNCHANGED)
3. Async dual-engine fetcher retry with exponential backoff & Playwright fallback
4. UTF-8-SIG BOM CSV & Styled Excel (.xlsx) export
5. BYOK LINE / Telegram webhook notification formatting
"""
from __future__ import annotations
import asyncio
from pathlib import Path
from typing import Dict, List

import openpyxl
import pytest

from scraper_pipeline import (
    AsyncDualEngineFetcher,
    DiffSnapshotStore,
    DiffStatus,
    DirectoryContactRecord,
    RetryConfig,
    export_diff_to_csv,
    export_diff_to_excel,
    normalize_fullwidth_text,
    parse_and_normalize_phone,
    send_webhook_alert,
)


def test_fullwidth_and_phone_extension_normalization() -> None:
    assert normalize_fullwidth_text("（０３）５２１－６１２１") == "(03)521-6121"
    phone, ext = parse_and_normalize_phone("（０４）２２２８－９１１１ 分機 22100")
    assert phone == "(04)2228-9111"
    assert ext == "22100"

    phone2, ext2 = parse_and_normalize_phone("03-5518101#305")
    assert phone2 == "(03)551-8101"
    assert ext2 == "305"


def test_diff_snapshot_store_tracks_new_updated_and_unchanged(tmp_path: Path) -> None:
    store = DiffSnapshotStore(snapshot_path=tmp_path / "state.json")
    rec_a = DirectoryContactRecord.from_raw(
        record_id="R-01",
        region="新竹市",
        organization="新竹市政府 行政處",
        department_or_duty="文書科",
        phone="03-5216121#101",
        scraped_at="2026-10-28 09:00:00",
    )
    first_run = store.compare_and_update([rec_a])
    assert len(first_run) == 1
    assert first_run[0].status == DiffStatus.NEW

    # Second run: R-01 unchanged, R-02 new
    rec_b = DirectoryContactRecord.from_raw(
        record_id="R-02",
        region="苗栗縣",
        organization="苗栗縣政府 農業處",
        department_or_duty="農務科",
        phone="037-559123#200",
        scraped_at="2026-10-28 10:00:00",
    )
    second_run = store.compare_and_update([rec_a, rec_b])
    assert second_run[0].status == DiffStatus.UNCHANGED
    assert second_run[1].status == DiffStatus.NEW

    # Third run: R-01 phone extension updated
    rec_a_updated = DirectoryContactRecord.from_raw(
        record_id="R-01",
        region="新竹市",
        organization="新竹市政府 行政處",
        department_or_duty="文書科",
        phone="03-5216121#999",
        scraped_at="2026-10-28 11:00:00",
    )
    third_run = store.compare_and_update([rec_a_updated, rec_b])
    assert third_run[0].status == DiffStatus.UPDATED
    assert third_run[1].status == DiffStatus.UNCHANGED


def test_async_dual_engine_retry_and_playwright_fallback() -> None:
    attempts = {"static": 0, "pw": 0}

    async def flaky_static(url: str, ua: str) -> str:
        attempts["static"] += 1
        if attempts["static"] < 2:
            raise ConnectionError("Temporary network glitch")
        return "<html><body><h1>Directory Table Ready</h1><table><tr><td>Row 1</td></tr></table></body></html>"

    async def pw_fallback(url: str) -> str:
        attempts["pw"] += 1
        return "<html><body><div id='hydrated'>Dynamic SPA Table Content</div></body></html>"

    fetcher = AsyncDualEngineFetcher(
        config=RetryConfig(max_retries=3, base_delay_sec=0.005, jitter_max_sec=0.002),
        static_transport=flaky_static,
        playwright_fallback=pw_fallback,
    )
    html = asyncio.run(fetcher.fetch_html("https://example.gov.tw/list"))
    assert "Directory Table Ready" in html
    assert attempts["static"] == 2
    assert attempts["pw"] == 0

    # Force JS fallback
    spa_html = asyncio.run(fetcher.fetch_html("https://spa.example.gov.tw", require_js=True))
    assert "Dynamic SPA Table Content" in spa_html
    assert attempts["pw"] == 1


def test_csv_and_excel_exporters_and_webhook_notifier(tmp_path: Path) -> None:
    store = DiffSnapshotStore()
    rec = DirectoryContactRecord.from_raw(
        record_id="R-10",
        region="臺中市",
        organization="臺中市政府 交通局",
        department_or_duty="運輸管理科",
        phone="04-22289111#60300",
        scraped_at="2026-10-28 10:00:00",
    )
    diffs = store.compare_and_update([rec])

    csv_file = export_diff_to_csv(diffs, tmp_path / "out.csv")
    xlsx_file = export_diff_to_excel(diffs, tmp_path / "out.xlsx")

    raw_bytes = csv_file.read_bytes()
    assert raw_bytes.startswith(b"\xef\xbb\xbf")  # UTF-8 BOM check
    assert "臺中市政府 交通局" in csv_file.read_text(encoding="utf-8-sig")

    wb = openpyxl.load_workbook(xlsx_file)
    ws = wb.active
    assert ws.cell(row=2, column=1).value == "NEW"
    assert ws.cell(row=2, column=4).value == "臺中市政府 交通局"

    sent_payloads: List[Dict[str, object]] = []
    summary = send_webhook_alert(
        "Test Pipeline",
        diffs,
        channel="line",
        sender=lambda ch, p: sent_payloads.append(p) or True,
    )
    assert summary.should_notify is True
    assert summary.new_count == 1
    assert len(sent_payloads) == 1
